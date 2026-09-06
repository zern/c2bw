from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"D:\pythonCode\c2bw\自愿退学申请书.docx")


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=11, bold=None):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def set_paragraph_layout(paragraph, *, first_line=0, left=0, before=0, after=0,
                         line=16, keep_with_next=False, alignment=None):
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(first_line)
    fmt.left_indent = Pt(left)
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = Pt(line)
    fmt.keep_with_next = keep_with_next
    fmt.widow_control = False
    if alignment is not None:
        paragraph.alignment = alignment


def add_text_paragraph(doc, text, *, bold=False, east_asia="宋体", size=11,
                       first_line=0, left=0, before=0, after=0, line=16,
                       keep_with_next=False, alignment=None):
    p = doc.add_paragraph()
    set_paragraph_layout(
        p,
        first_line=first_line,
        left=left,
        before=before,
        after=after,
        line=line,
        keep_with_next=keep_with_next,
        alignment=alignment,
    )
    run = p.add_run(text)
    set_run_font(run, east_asia=east_asia, size=size, bold=bold)
    return p


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.PORTRAIT
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.58)
section.bottom_margin = Inches(0.52)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)
section.header_distance = Inches(0.2)
section.footer_distance = Inches(0.2)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

title_style = doc.styles["Title"]
title_style.font.name = "Times New Roman"
title_style.font.size = Pt(19)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

title = doc.add_paragraph(style="Title")
set_paragraph_layout(
    title,
    before=0,
    after=9,
    line=22,
    keep_with_next=True,
    alignment=WD_ALIGN_PARAGRAPH.CENTER,
)
title_run = title.add_run("自愿退学申请书")
set_run_font(title_run, east_asia="黑体", size=19, bold=True)

add_text_paragraph(
    doc,
    "尊敬的学校领导、班主任老师：",
    after=2,
    line=16,
    keep_with_next=True,
)

intro = doc.add_paragraph()
set_paragraph_layout(intro, first_line=22, after=2, line=16)
run = intro.add_run("本人系本校高__年级__班学生")
set_run_font(run)
run = intro.add_run("________。经过本人慎重思考，并与父母（法定监护人）充分沟通、协商一致后，自愿主动向学校申请退学")
set_run_font(run, bold=True)
run = intro.add_run("。")
set_run_font(run)

add_text_paragraph(
    doc,
    "本人退学纯属个人自愿行为，与学校、老师无关。本人已充分知晓并自愿承担以下全部责任与后果：",
    first_line=22,
    after=2,
    line=16,
)

add_text_paragraph(
    doc,
    "1. 学籍后果知情",
    bold=True,
    east_asia="黑体",
    before=1,
    after=0,
    line=16,
    keep_with_next=True,
)
add_text_paragraph(
    doc,
    "本人清楚退学一经学校审批、上报教育局备案后，学籍将作退学处理，不再保留本校在读学籍，由此造成无法复学、无法参加高考、无法继续就读本校等一切学籍后果，均由本人及家庭自行承担。",
    first_line=22,
    after=1,
    line=16,
)

add_text_paragraph(
    doc,
    "2. 安全责任自负",
    bold=True,
    east_asia="黑体",
    before=1,
    after=0,
    line=16,
    keep_with_next=True,
)
add_text_paragraph(
    doc,
    "自退学申请批准、离校之日起，本人不再属于在校学生。离校后的一切人身安全、出行安全、生活、务工、社会活动等所有风险与意外事故，全部由本人及监护人自行负责，与学校无任何关系，学校不承担任何后续安全及管理责任。",
    first_line=22,
    after=1,
    line=16,
)

add_text_paragraph(
    doc,
    "3. 自愿放弃在校就读权利",
    bold=True,
    east_asia="黑体",
    before=1,
    after=0,
    line=16,
    keep_with_next=True,
)
add_text_paragraph(
    doc,
    "本人自愿放弃剩余学业学习、学校管理、各类评优及考试资格，绝不因退学后续问题追究学校任何责任。",
    first_line=22,
    after=2,
    line=16,
)

add_text_paragraph(
    doc,
    "本人及家长承诺：本次退学完全自愿、无胁迫、无纠纷，所有决定真实有效。",
    first_line=22,
    after=1,
    line=16,
)
add_text_paragraph(
    doc,
    "恳请学校领导予以批准！",
    first_line=22,
    after=4,
    line=16,
)

signature_left = 230
add_text_paragraph(
    doc,
    "申请人（学生签名）：__________",
    left=signature_left,
    after=1,
    line=16,
)
add_text_paragraph(
    doc,
    "家长（监护人签名）：__________",
    left=signature_left,
    after=1,
    line=16,
)
add_text_paragraph(
    doc,
    "家长联系电话：__________",
    left=signature_left,
    after=1,
    line=16,
)
add_text_paragraph(
    doc,
    "日期：______年____月____日",
    left=signature_left,
    after=3,
    line=16,
)

add_text_paragraph(
    doc,
    "特点（完全符合学校要求）",
    bold=True,
    east_asia="黑体",
    size=10.5,
    before=1,
    after=0,
    line=15,
)

expected = [
    "自愿退学申请书",
    "尊敬的学校领导、班主任老师：",
    "本人系本校高__年级__班学生________。经过本人慎重思考，并与父母（法定监护人）充分沟通、协商一致后，自愿主动向学校申请退学。",
    "本人退学纯属个人自愿行为，与学校、老师无关。本人已充分知晓并自愿承担以下全部责任与后果：",
    "1. 学籍后果知情",
    "本人清楚退学一经学校审批、上报教育局备案后，学籍将作退学处理，不再保留本校在读学籍，由此造成无法复学、无法参加高考、无法继续就读本校等一切学籍后果，均由本人及家庭自行承担。",
    "2. 安全责任自负",
    "自退学申请批准、离校之日起，本人不再属于在校学生。离校后的一切人身安全、出行安全、生活、务工、社会活动等所有风险与意外事故，全部由本人及监护人自行负责，与学校无任何关系，学校不承担任何后续安全及管理责任。",
    "3. 自愿放弃在校就读权利",
    "本人自愿放弃剩余学业学习、学校管理、各类评优及考试资格，绝不因退学后续问题追究学校任何责任。",
    "本人及家长承诺：本次退学完全自愿、无胁迫、无纠纷，所有决定真实有效。",
    "恳请学校领导予以批准！",
    "申请人（学生签名）：__________",
    "家长（监护人签名）：__________",
    "家长联系电话：__________",
    "日期：______年____月____日",
    "特点（完全符合学校要求）",
]

actual = [p.text for p in doc.paragraphs]
assert actual == expected, (actual, expected)

doc.core_properties.title = "自愿退学申请书"
doc.core_properties.subject = "学生自愿退学申请"
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""
doc.save(OUTPUT)
print(OUTPUT)
